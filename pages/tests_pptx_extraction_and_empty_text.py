"""
D9: PowerPoint decks are read correctly, and no analysis runs on empty text.

What was broken:

- confirm_analyze_founder_profile (the investor's paid "Analyze" step) sent
  every deck through the PDF extractor, including .pptx files. PyPDF2 fails
  on a PowerPoint file, the failure was swallowed, and the deck was analyzed
  from nothing -- with the investor charged a credit.
- Nothing downstream refused empty text. The chunker returns no chunks and
  the pipeline still called Claude, producing a memo from zero content.
- Once any pitch-deck document existed -- even an empty or failed one -- the
  analyze and confirm steps returned it as "ready" forever, so an affected
  founder could never be re-analyzed.

The invariant: no usable extracted text -> no document, no analysis charge,
no analysis task, no Claude call. "Usable" means something other than
whitespace (a deck of empty text boxes extracts to newlines).

Out of scope, deliberately: OCR for image-only decks (they get a clear
message instead), .ppt/.pptm, and any cleanup of existing broken rows -- the
repaired confirm step re-extracts those the next time an investor asks.
"""
import io
import tempfile
from datetime import timedelta
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from pptx import Presentation
from pptx.util import Inches

from matchmaking.models import Application, InvestorApplication
from matchmaking.tests import _mock_embedding_generation
from zelda_api.models import AnalysisCreditCharge
from zelda_api.vector_models import DocumentSource, IntelligenceMemo

User = get_user_model()

SLIDE_TEXT = 'We raised $2M ARR growing 20% month over month'


def pptx_bytes(slides):
    """A real .pptx. Each slide is 'text' (with SLIDE_TEXT), 'image' (no text shapes) or 'empty' (an empty text box)."""
    prs = Presentation()
    for kind in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
        if kind == 'text':
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = SLIDE_TEXT
        elif kind == 'empty':
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def refuse(*args, **kwargs):
    raise AssertionError('this must not be called')


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class _Deck(TestCase):

    def setUp(self):
        _mock_embedding_generation(self)
        self.founder_user = User.objects.create_user('d9_founder', password='x')
        self.founder = Application.objects.create(
            user=self.founder_user, company_name='Acme Decks', founder_name='F', email='f@t.com',
            description='test', sector='SaaS', stage='Seed',
        )
        self.investor_user = User.objects.create_user('d9_investor', password='x')
        InvestorApplication.objects.create(user=self.investor_user)
        self.pipeline = self._patch('zelda_api.tasks.process_document_pipeline.delay')

    def _patch(self, target, **kwargs):
        patcher = mock.patch(target, **kwargs)
        started = patcher.start()
        self.addCleanup(patcher.stop)
        return started

    def _deck(self, content, name='deck.pptx'):
        self.founder.pitch_deck = SimpleUploadedFile(name, content)
        self.founder.save()

    def _confirm(self):
        self.client.force_login(self.investor_user)
        return self.client.post(reverse('zelda_api:analyze_founder_confirm', args=[self.founder_user.username]))

    def _analyze(self):
        self.client.force_login(self.investor_user)
        return self.client.get(reverse('zelda_api:analyze_founder', args=[self.founder_user.username]))

    def _existing(self, text='', status='analyzed', memo=False, minutes_ago=60):
        doc = DocumentSource.objects.create(
            uploaded_by=self.founder_user, filename='decks/old.pptx', source_entity='Acme Decks',
            document_type='pitch_deck', status=status, raw_text_preview=text[:1000], raw_text_full=text,
        )
        DocumentSource.objects.filter(pk=doc.pk).update(created_at=timezone.now() - timedelta(minutes=minutes_ago))
        if memo:
            IntelligenceMemo.objects.create(
                document=doc, executive_summary='Not disclosed in pitch deck.', business_model_analysis='Not disclosed.',
                evidence_level='PARTLY_EVIDENCED', completeness_score=0.1, citations_count=0,
            )
        doc.refresh_from_db()
        return doc

    def assert_nothing_spent(self):
        self.assertFalse(DocumentSource.objects.filter(uploaded_by=self.founder_user, created_at__gte=self.started).exists())
        self.assertFalse(AnalysisCreditCharge.objects.exists())
        self.pipeline.assert_not_called()


class PowerPointExtractionTests(_Deck):

    def test_a_powerpoint_deck_is_read_with_the_powerpoint_extractor(self):
        self._deck(pptx_bytes(['text', 'text', 'image']))
        response = self._confirm()
        self.assertEqual(response.status_code, 200, response.content[:300])
        self.assertEqual(response.json()['status'], 'processing')
        doc = DocumentSource.objects.get(id=response.json()['document_id'])
        self.assertIn(SLIDE_TEXT, doc.raw_text_full)
        self.assertEqual(doc.total_pages, 3)
        self.pipeline.assert_called_once()
        self.assertIn(SLIDE_TEXT, self.pipeline.call_args.args[1])
        self.assertTrue(AnalysisCreditCharge.objects.filter(document=doc, user=self.investor_user).exists())

    def test_the_pdf_extractor_is_never_used_for_a_powerpoint_deck(self):
        self._deck(pptx_bytes(['text']))
        with mock.patch('zelda_api.utils._extract_pdf_text', side_effect=refuse):
            response = self._confirm()
        self.assertEqual(response.status_code, 200, response.content[:300])

    def test_a_pdf_deck_still_goes_through_the_pdf_extractor(self):
        self._deck(b'%PDF-1.4 fake', name='deck.pdf')
        with mock.patch('zelda_api.utils._extract_pdf_text', return_value=('Deck text from a PDF.', 4)), \
                mock.patch('zelda_api.utils._extract_pptx_text', side_effect=refuse):
            response = self._confirm()
        self.assertEqual(response.status_code, 200, response.content[:300])
        doc = DocumentSource.objects.get(id=response.json()['document_id'])
        self.assertEqual(doc.raw_text_full, 'Deck text from a PDF.')
        self.assertEqual(doc.total_pages, 4)

    def test_the_deck_is_read_through_storage_not_a_local_file_path(self):
        # Uploads move to S3 at launch, where a FieldFile has no local .path.
        from django.db.models.fields.files import FieldFile

        self._deck(pptx_bytes(['text']))
        with mock.patch.object(FieldFile, 'path', new_callable=mock.PropertyMock,
                               side_effect=NotImplementedError('no local path')):
            response = self._confirm()
        self.assertEqual(response.status_code, 200, response.content[:300])

    def test_the_shared_dispatcher_reads_a_stored_powerpoint(self):
        from zelda_api.utils import extract_text_from_file

        self._deck(pptx_bytes(['text', 'image']))
        with self.founder.pitch_deck.open('rb') as deck:
            text, pages = extract_text_from_file(deck)
        self.assertIn(SLIDE_TEXT, text)
        self.assertEqual(pages, 2)


class UnreadableDeckTests(_Deck):

    def setUp(self):
        super().setUp()
        self.started = timezone.now()

    def _assert_refused(self, response):
        from zelda_api.utils import UNREADABLE_DOCUMENT_MESSAGE

        self.assertEqual(response.status_code, 422, response.content[:300])
        body = response.json()
        self.assertEqual(body['status'], 'error')
        self.assertEqual(body['message'], UNREADABLE_DOCUMENT_MESSAGE)
        self.assert_nothing_spent()

    def test_the_message_says_what_happened_and_what_to_do(self):
        from zelda_api.utils import UNREADABLE_DOCUMENT_MESSAGE

        self.assertIn('readable text', UNREADABLE_DOCUMENT_MESSAGE)
        self.assertIn('images', UNREADABLE_DOCUMENT_MESSAGE)

    def test_an_image_only_deck_is_refused_and_nothing_is_spent(self):
        self._deck(pptx_bytes(['image', 'image']))
        self._assert_refused(self._confirm())

    def test_a_deck_of_empty_text_boxes_counts_as_unreadable(self):
        self._deck(pptx_bytes(['empty', 'empty']))
        self._assert_refused(self._confirm())

    def test_a_file_that_cannot_be_opened_is_refused_the_same_way(self):
        self._deck(b'this is not a presentation')
        self._assert_refused(self._confirm())

    def test_an_unreadable_pdf_is_refused_the_same_way(self):
        self._deck(b'%PDF-1.4 fake', name='deck.pdf')
        with mock.patch('zelda_api.utils._extract_pdf_text', return_value=('   \n', 2)):
            self._assert_refused(self._confirm())

    def test_no_claude_client_is_created_for_an_unreadable_deck(self):
        self._deck(pptx_bytes(['image']))
        with mock.patch('zelda_api.anthropic_client.background_anthropic_client', side_effect=refuse) as claude:
            self._assert_refused(self._confirm())
        claude.assert_not_called()


class RecoveryTests(_Deck):
    """A broken existing document is never offered as ready; the deck is read again."""

    def setUp(self):
        super().setUp()
        self._deck(pptx_bytes(['text']))

    def test_an_existing_document_with_no_text_is_not_offered_as_ready(self):
        self._existing(text='', status='ingested')  # founder2's #993
        self.assertEqual(self._analyze().json()['status'], 'confirm_required')

    def test_a_hollow_analyzed_document_is_not_offered_as_ready(self):
        self._existing(text='\n\n', status='analyzed', memo=True)
        self.assertEqual(self._analyze().json()['status'], 'confirm_required')

    def test_an_errored_document_is_not_offered_as_ready(self):
        self._existing(text='Real deck text.', status='error')
        self.assertEqual(self._analyze().json()['status'], 'confirm_required')

    def test_confirming_reads_the_deck_again_instead_of_returning_the_broken_document(self):
        broken = self._existing(text='', status='ingested')
        response = self._confirm()
        self.assertEqual(response.status_code, 200, response.content[:300])
        self.assertEqual(response.json()['status'], 'processing')
        fresh = DocumentSource.objects.get(id=response.json()['document_id'])
        self.assertNotEqual(fresh.id, broken.id)
        self.assertIn(SLIDE_TEXT, fresh.raw_text_full)
        self.assertTrue(AnalysisCreditCharge.objects.filter(document=fresh, user=self.investor_user).exists())

    def test_after_recovery_the_fresh_document_is_the_one_offered_everywhere(self):
        from zelda_api.ic_memo import latest_analyzed_pitch_deck_and_memo

        hollow = self._existing(text='', status='analyzed', memo=True)
        fresh_id = self._confirm().json()['document_id']
        fresh = DocumentSource.objects.get(id=fresh_id)
        fresh.status = 'analyzed'
        fresh.save()
        IntelligenceMemo.objects.create(
            document=fresh, executive_summary='A real summary.', business_model_analysis='Real thesis.',
            evidence_level='PARTLY_EVIDENCED', completeness_score=0.7, citations_count=3,
        )

        ready = self._analyze().json()
        self.assertEqual(ready['status'], 'ready')
        self.assertEqual(ready['document_id'], fresh.id)
        doc, memo = latest_analyzed_pitch_deck_and_memo(self.founder_user)
        self.assertEqual(doc.id, fresh.id)
        self.assertNotEqual(doc.id, hollow.id)

    def test_a_usable_document_is_still_reused_for_free(self):
        good = self._existing(text='A real deck.', status='analyzed', memo=True)
        ready = self._analyze().json()
        self.assertEqual((ready['status'], ready['document_id']), ('ready', good.id))
        response = self._confirm()
        self.assertEqual(response.json()['status'], 'ready')
        self.assertEqual(response.json()['document_id'], good.id)
        self.assertFalse(AnalysisCreditCharge.objects.exists())
        self.pipeline.assert_not_called()

    def test_a_usable_document_still_processing_is_not_charged_again(self):
        in_flight = self._existing(text='A real deck.', status='chunking')
        response = self._confirm()
        self.assertEqual(response.json()['status'], 'ready')
        self.assertEqual(response.json()['document_id'], in_flight.id)
        self.assertEqual(DocumentSource.objects.filter(uploaded_by=self.founder_user).count(), 1)
        self.assertFalse(AnalysisCreditCharge.objects.exists())


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class OwnerUploadTests(TestCase):
    """DocumentIngestView already chose the right extractor; it let whitespace through."""

    def setUp(self):
        _mock_embedding_generation(self)
        self.user = User.objects.create_user('d9_owner', password='x')
        self.client.force_login(self.user)
        patcher = mock.patch('zelda_api.pipeline_views.process_document_pipeline.delay')
        self.pipeline = patcher.start()
        self.addCleanup(patcher.stop)

    def _upload(self, content):
        return self.client.post(reverse('zelda_api:document_ingest'), {
            'file': SimpleUploadedFile('deck.pptx', content), 'document_type': 'pitch_deck', 'source_entity': 'Acme',
        })

    def test_an_uploaded_powerpoint_is_read(self):
        response = self._upload(pptx_bytes(['text']))
        self.assertEqual(response.status_code, 201, response.content[:300])
        self.assertIn(SLIDE_TEXT, DocumentSource.objects.get(id=response.json()['document_id']).raw_text_full)

    def test_a_deck_of_empty_text_boxes_is_refused_and_not_queued(self):
        from zelda_api.utils import UNREADABLE_DOCUMENT_MESSAGE

        response = self._upload(pptx_bytes(['empty', 'empty']))
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()['error'], UNREADABLE_DOCUMENT_MESSAGE)
        self.assertFalse(DocumentSource.objects.exists())
        self.pipeline.assert_not_called()

    def test_an_image_only_deck_is_refused(self):
        response = self._upload(pptx_bytes(['image']))
        self.assertEqual(response.status_code, 400)
        self.assertFalse(DocumentSource.objects.exists())


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class PipelineGuardTests(TestCase):
    """Belt and braces: whatever queued it, the pipeline never analyzes empty text."""

    def setUp(self):
        _mock_embedding_generation(self)
        self.user = User.objects.create_user('d9_pipeline', password='x')
        self.doc = DocumentSource.objects.create(
            uploaded_by=self.user, filename='deck.pptx', source_entity='Acme', document_type='pitch_deck',
            status='ingested',
        )
        # A Mock that raises still records being called. The pipeline turns
        # exceptions into an error result, so refusing alone would hide an
        # attempted call; the tests assert the client was never requested.
        self.claude = mock.Mock(side_effect=refuse)
        patcher = mock.patch('zelda_api.anthropic_client.background_anthropic_client', new=self.claude)
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_the_fundraising_pipeline_refuses_empty_text_without_calling_claude(self):
        from zelda_api.intelligence_pipeline import intelligence_pipeline

        for text in ('', '\n\n', '   '):
            with self.subTest(text=text):
                result = intelligence_pipeline.process_document(self.doc, text)
                self.assertEqual(result['status'], 'error')
                self.doc.refresh_from_db()
                self.assertEqual(self.doc.status, 'error')
                self.assertFalse(hasattr(self.doc, 'memo'))
                self.assertFalse(self.doc.chunks.exists())
                self.claude.assert_not_called()

    def test_the_valuation_pipeline_refuses_empty_text_too(self):
        from zelda_api.intelligence_pipeline import intelligence_pipeline

        self.doc.document_type = 'business_valuation'
        self.doc.save()
        result = intelligence_pipeline.process_valuation_document(self.doc, '\n')
        self.assertEqual(result['status'], 'error')
        self.doc.refresh_from_db()
        self.assertEqual(self.doc.status, 'error')
        self.claude.assert_not_called()

    def test_the_pipeline_task_does_not_retry_or_log_an_unreadable_document(self):
        from ops.models import FailedTaskLog
        from zelda_api.tasks import process_document_pipeline

        # Retrying can't make an empty deck readable, and a retry would reach
        # for the broker, so it must never be attempted.
        with mock.patch.object(process_document_pipeline, 'retry', side_effect=refuse) as retry:
            result = process_document_pipeline.run(self.doc.id, '')
        retry.assert_not_called()
        self.assertEqual(result['status'], 'error')
        self.assertFalse(FailedTaskLog.objects.exists())
        self.doc.refresh_from_db()
        self.assertEqual(self.doc.status, 'error')


class UsableTextTests(TestCase):

    def test_whitespace_is_not_usable_text(self):
        from zelda_api.utils import has_usable_text

        for text, usable in (('', False), (None, False), ('\n\n', False), (' \t ', False), ('x', True)):
            with self.subTest(text=text):
                self.assertIs(has_usable_text(text), usable)
