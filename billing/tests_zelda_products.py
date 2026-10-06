import json
from unittest import mock
from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, SimpleTestCase, RequestFactory, override_settings
from django.urls import reverse
from django.utils import timezone
from .models import ZeldaOrder
from .zelda_catalog import PRODUCTS, selected_reports
from zelda_api.vector_models import DocumentSource, IntelligenceMemo, BusinessValuationReport

User = get_user_model()


@override_settings(STRIPE_SECRET_KEY='sk_test_zelda_ci_only')
class ZeldaProductTests(TestCase):
    def setUp(self):
        from matchmaking.tests import _mock_embedding_generation
        _mock_embedding_generation(self)
        cache.clear()
        self.addCleanup(cache.clear)
        from zelda_api.sec_company_identity import search_listed_companies
        self.real_listed_search = search_listed_companies
        listed = mock.patch('zelda_api.sec_company_identity.search_listed_companies', return_value=[])
        listed.start()
        self.addCleanup(listed.stop)
        self.user = User.objects.create_user('product_customer', password='x')
        self.other = User.objects.create_user('product_other', password='x')
        self.client.force_login(self.user)
        self.source = DocumentSource.objects.create(
            uploaded_by=self.user, source_entity='External Example Inc', filename='evidence.txt',
            raw_text_full='Company-provided research evidence describing the business and market. ' * 20,
            is_external_subject=True, is_product_input=True,
        )

    def order(self, **kwargs):
        defaults = dict(user=self.user, source_document=self.source, product='ic_memo', reports=['ic_memo'],
                        amount=499, stripe_session_id='cs_product_test')
        defaults.update(kwargs)
        return ZeldaOrder.objects.create(**defaults)

    def checkout(self, **kwargs):
        data = dict(document_id=self.source.id, product='ic_memo')
        data.update(kwargs)
        return self.client.post(reverse('billing:zelda_checkout'), json.dumps(data), content_type='application/json')

    def event(self, order, **kwargs):
        data = dict(id=order.stripe_session_id, mode='payment', status='complete', currency='usd', amount_total=order.amount, payment_status='paid',
                    metadata={'purpose':'zelda_product','order_id':str(order.id),'user_id':str(self.user.id)})
        data.update(kwargs)
        return data

    def test_catalog_matches_all_seven_names_and_prices(self):
        response = self.client.get(reverse('billing:zelda_catalog'))
        products = response.json()['products']
        self.assertEqual({p['name']:p['amount'] for p in products}, {
            'Zelda Intelligence Memo':199,
            'Zelda IC Memo':499,
            'Truth Delta Report':1999,
            'Zelda Complete Intelligence Bundle':4999,
            'Zelda 3-Report Pack':2499,
            'Entity Integrity Report':999,
            'Business valuation':99,
        })

    def test_pack_requires_three_distinct_reports_and_complete_bundle_has_five(self):
        self.assertEqual(len(selected_reports('complete_bundle')), 5)
        for reports in ([], ['ic_memo'], ['ic_memo']*3, ['ic_memo','entity','unknown']):
            with self.assertRaises(ValueError): selected_reports('three_pack', reports)
        self.assertEqual(len(selected_reports('three_pack', ['ic_memo','entity','truth_delta'])),3)

    @mock.patch('billing.zelda_views.stripe_price', return_value='price_ic')
    @mock.patch('stripe.checkout.Session.create', return_value={'id':'cs_created','url':'https://checkout.stripe.com/c/test'})
    def test_checkout_uses_catalog_price_and_reuses_session_without_starting_generation(self, checkout, price):
        with mock.patch('billing.tasks.fulfill_zelda_order.delay') as generate:
            self.assertEqual(self.checkout().status_code, 200)
            order = ZeldaOrder.objects.get()
            with mock.patch('stripe.checkout.Session.retrieve', return_value=self.event(
                    order, status='open', payment_status='unpaid', url='https://checkout.stripe.com/c/test')):
                self.assertEqual(self.checkout().status_code, 200)
        checkout.assert_called_once()
        self.assertEqual(checkout.call_args.kwargs['line_items'], [{'price':'price_ic','quantity':1}])
        self.assertEqual(checkout.call_args.kwargs['mode'],'payment')
        self.assertEqual(checkout.call_args.kwargs['api_key'],'sk_test_zelda_ci_only')
        generate.assert_not_called()
        self.assertEqual(ZeldaOrder.objects.get().amount,499)

    @mock.patch('billing.zelda_views.stripe_price', return_value='price_entity')
    @mock.patch('stripe.checkout.Session.create', return_value={'id':'cs_replacement','url':'https://checkout.stripe.com/c/new'})
    def test_expired_entity_checkout_can_resume_with_same_uploaded_evidence(self, create, price):
        old = self.order(product='entity', reports=['entity'], amount=999,
                         checkout_url='https://checkout.stripe.com/c/expired')
        with mock.patch('stripe.checkout.Session.retrieve', return_value=self.event(
                old, status='expired', payment_status='unpaid')):
            response = self.checkout(product='entity')
        self.assertEqual(response.json()['checkout_url'], 'https://checkout.stripe.com/c/new')
        old.refresh_from_db()
        self.assertEqual(old.status, 'canceled')
        replacement = ZeldaOrder.objects.exclude(pk=old.pk).get()
        self.assertEqual(replacement.source_document_id, self.source.pk)
        self.assertEqual(create.call_args.kwargs['idempotency_key'], f'zelda-order-{replacement.pk}')

    def test_complete_checkout_does_not_start_another_payment(self):
        for payment_status in ('paid', 'unpaid'):
            with self.subTest(payment_status=payment_status):
                ZeldaOrder.objects.all().delete()
                order = self.order(checkout_url='https://checkout.stripe.com/c/old')
                with mock.patch('stripe.checkout.Session.retrieve', return_value=self.event(
                        order, payment_status=payment_status)), mock.patch('stripe.checkout.Session.create') as create, \
                        mock.patch('billing.tasks.fulfill_zelda_order.delay'):
                    response = self.checkout()
                self.assertEqual(response.json()['order_url'], reverse('billing:zelda_order', args=[order.pk]))
                create.assert_not_called()
                order.refresh_from_db()
                self.assertEqual(order.status, 'paid' if payment_status == 'paid' else 'awaiting_payment')

    def test_checkout_refuses_unverified_or_mismatched_saved_session(self):
        order = self.order(checkout_url='https://checkout.stripe.com/c/old')
        for session in (self.event(order, id='cs_other', status='expired', payment_status='unpaid'),
                        self.event(order, amount_total=1, status='expired', payment_status='unpaid')):
            with mock.patch('stripe.checkout.Session.retrieve', return_value=session), \
                    mock.patch('stripe.checkout.Session.create') as create:
                self.assertEqual(self.checkout().status_code, 503)
                create.assert_not_called()
        with mock.patch('stripe.checkout.Session.retrieve', side_effect=RuntimeError('Stripe unavailable')), \
                mock.patch('stripe.checkout.Session.create') as create:
            self.assertEqual(self.checkout().status_code, 503)
            create.assert_not_called()
        order.refresh_from_db()
        self.assertEqual(order.status, 'awaiting_payment')

    def test_order_has_checkout_action_even_if_session_creation_failed(self):
        order = self.order(stripe_session_id=None)
        page = self.client.get(reverse('billing:zelda_order', args=[order.pk]))
        self.assertContains(page, 'id="order-checkout"')
        self.assertContains(page, reverse('billing:zelda_checkout'))

    def test_checkout_refuses_other_users_evidence_and_invalid_pack(self):
        self.client.force_login(self.other)
        self.assertEqual(self.checkout().status_code,404)
        self.client.force_login(self.user)
        self.assertEqual(self.checkout(product='three_pack',reports=['entity']).status_code,400)
        self.assertFalse(ZeldaOrder.objects.exists())

    @mock.patch('billing.zelda_views.stripe_price', side_effect=ValueError('catalog not configured'))
    def test_catalog_failure_does_not_charge_or_generate(self, price):
        with mock.patch('stripe.checkout.Session.create') as checkout, mock.patch('billing.tasks.fulfill_zelda_order.delay') as generate:
            self.assertEqual(self.checkout().status_code,503)
        checkout.assert_not_called();generate.assert_not_called()
        self.assertEqual(ZeldaOrder.objects.get().status,'awaiting_payment')

    def test_unpaid_webhook_cannot_fulfill_and_paid_webhook_can(self):
        from .zelda_views import handle_product_event
        order=self.order()
        with mock.patch('billing.tasks.fulfill_zelda_order.delay') as generate:
            with self.captureOnCommitCallbacks(execute=True):
                handle_product_event('checkout.session.completed',self.event(order,payment_status='unpaid'))
            generate.assert_not_called()
            with self.captureOnCommitCallbacks(execute=True):
                handle_product_event('checkout.session.async_payment_succeeded',self.event(order))
            generate.assert_called_once_with(str(order.id))
        order.refresh_from_db();self.assertEqual(order.status,'paid');self.assertIsNotNone(order.paid_at)

    def test_signed_webhook_dispatches_product_event(self):
        order=self.order()
        event={'type':'checkout.session.completed','data':{'object':self.event(order)}}
        with mock.patch('stripe.Webhook.construct_event',return_value=event), mock.patch('billing.tasks.fulfill_zelda_order.delay') as generate:
            with self.captureOnCommitCallbacks(execute=True):
                response=self.client.post(reverse('billing:stripe_webhook'),data='{}',content_type='application/json')
        self.assertEqual(response.status_code,200);generate.assert_called_once()

    def test_amount_currency_user_and_order_mismatch_are_refused(self):
        from .zelda_views import handle_product_event
        order=self.order()
        for override in ({'amount_total':1},{'currency':'eur'},{'metadata':{'order_id':str(order.id),'user_id':str(self.other.id)}}, {'metadata':{'order_id':'wrong'}}):
            with self.assertRaises(ValueError): handle_product_event('checkout.session.completed',self.event(order,**override))
        order.refresh_from_db();self.assertEqual(order.status,'awaiting_payment')

    def test_duplicate_paid_events_cannot_create_duplicate_worker_outputs(self):
        from .tasks import fulfill_zelda_order
        order=self.order(status='paid',paid_at=timezone.now())
        with mock.patch('zelda_api.tasks.process_document_pipeline.delay') as pipeline, mock.patch('billing.tasks.monitor_zelda_order.delay'):
            fulfill_zelda_order.run(str(order.id));fulfill_zelda_order.run(str(order.id))
        pipeline.assert_called_once()
        order.refresh_from_db();self.assertEqual(order.status,'processing')
        self.assertEqual(order.analysis_document.uploaded_by,self.user)
        self.assertTrue(order.analysis_document.is_external_subject)

    def test_payment_redirect_and_other_users_do_not_unlock_reports(self):
        order=self.order()
        self.assertEqual(self.client.get(reverse('billing:zelda_order',args=[order.id])).status_code,200)
        self.assertEqual(self.client.get(reverse('billing:zelda_report',args=[order.id,'ic_memo'])).status_code,404)
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(reverse('billing:zelda_order_status',args=[order.id])).status_code,404)

    @mock.patch('stripe.checkout.Session.retrieve')
    @mock.patch('billing.tasks.fulfill_zelda_order.delay')
    def test_paid_business_valuation_recovers_without_webhook_or_new_charge(self, generate, retrieve):
        order = self.order(product='valuation', reports=['valuation'], amount=99)
        retrieve.return_value = self.event(order)
        with self.captureOnCommitCallbacks(execute=True), mock.patch('stripe.checkout.Session.create') as create:
            response = self.client.get(reverse('billing:zelda_order_status', args=[order.id]))
        self.assertEqual(response.json()['status'], 'paid')
        self.assertEqual(response['Cache-Control'], 'no-store')
        order.refresh_from_db(); self.assertIsNotNone(order.paid_at)
        retrieve.assert_called_once_with(order.stripe_session_id, api_key='sk_test_zelda_ci_only')
        generate.assert_called_once_with(str(order.id)); create.assert_not_called()
        self.client.get(reverse('billing:zelda_order_status', args=[order.id]))
        retrieve.assert_called_once(); generate.assert_called_once()

    @mock.patch('stripe.checkout.Session.retrieve')
    @mock.patch('billing.tasks.fulfill_zelda_order.delay')
    def test_unpaid_open_pending_and_expired_checkout_never_unlock(self, generate, retrieve):
        order = self.order()
        for checkout_state in ('open', 'complete', 'expired'):
            cache.clear()
            retrieve.return_value = self.event(order, status=checkout_state, payment_status='unpaid')
            response = self.client.get(reverse('billing:zelda_order_status', args=[order.id]))
            self.assertEqual(response.json()['checkout_state'], checkout_state)
            self.assertEqual(response.json()['status'], 'canceled' if checkout_state == 'expired' else 'awaiting_payment')
            self.assertEqual(self.client.get(reverse('billing:zelda_report',args=[order.id,'ic_memo'])).status_code,404)
        generate.assert_not_called()

    @mock.patch('stripe.checkout.Session.retrieve')
    @mock.patch('billing.tasks.fulfill_zelda_order.delay')
    def test_payment_recovery_refuses_wrong_session_and_payment_binding(self, generate, retrieve):
        order = self.order()
        for override in ({'id':'cs_other'}, {'mode':'subscription'}, {'amount_total':1}, {'currency':'eur'},
                         {'status':'open'}, {'metadata':{'purpose':'subscription','order_id':str(order.id),'user_id':str(self.user.id)}},
                         {'metadata':{'purpose':'zelda_product','order_id':str(order.id),'user_id':str(self.other.id)}},
                         {'metadata':{'purpose':'zelda_product','order_id':'wrong','user_id':str(self.user.id)}}):
            cache.clear(); retrieve.return_value = self.event(order, **override)
            with self.assertLogs('billing.zelda_views', level='ERROR'):
                response = self.client.get(reverse('billing:zelda_order_status',args=[order.id]))
            self.assertEqual(response.json()['status'], 'awaiting_payment')
            self.assertTrue(response.json()['payment_check_unavailable'])
        generate.assert_not_called()

    @mock.patch('stripe.checkout.Session.retrieve', side_effect=RuntimeError('Stripe unavailable'))
    def test_failed_payment_check_is_throttled_and_owner_only(self, retrieve):
        order = self.order()
        url = reverse('billing:zelda_order_status',args=[order.id])
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(url).status_code,404); retrieve.assert_not_called()
        self.client.force_login(self.user)
        with self.assertLogs('billing.zelda_views', level='ERROR'):
            self.assertTrue(self.client.get(url).json()['payment_check_unavailable'])
        self.assertTrue(self.client.get(url).json()['payment_check_unavailable'])
        retrieve.assert_called_once()
        order.refresh_from_db(); self.assertIsNone(order.paid_at)

    @mock.patch('stripe.checkout.Session.retrieve')
    @mock.patch('billing.tasks.fulfill_zelda_order.delay')
    def test_queue_delivery_failure_retains_payment_and_can_recover(self, generate, retrieve):
        order = self.order()
        retrieve.return_value = self.event(order)
        generate.side_effect = RuntimeError('Broker unavailable')
        # TestCase delays on_commit beyond the request. Run it explicitly: a
        # delivery failure must not undo the recorded payment.
        with self.assertRaises(RuntimeError), self.captureOnCommitCallbacks(execute=True):
            response = self.client.get(reverse('billing:zelda_order_status',args=[order.id]))
        self.assertEqual(response.json()['status'],'paid')
        order.refresh_from_db(); self.assertIsNotNone(order.paid_at)
        cache.clear()
        # Already-paid recovery registers delivery outside an atomic block.
        with mock.patch('billing.zelda_views.transaction.on_commit', side_effect=lambda callback: callback()):
            with self.assertLogs('billing.zelda_views', level='ERROR'):
                response = self.client.get(reverse('billing:zelda_order_status',args=[order.id]))
            self.assertEqual(response.json()['status'],'paid')
            order.refresh_from_db(); self.assertIsNotNone(order.paid_at)
            cache.clear(); generate.side_effect = None
            self.client.get(reverse('billing:zelda_order_status',args=[order.id]))
        self.assertEqual(generate.call_count,3); retrieve.assert_called_once()

    @mock.patch('stripe.checkout.Session.retrieve')
    def test_processing_and_ready_orders_do_not_recheck_stripe(self, retrieve):
        order = self.order(status='processing', paid_at=timezone.now())
        self.client.get(reverse('billing:zelda_order_status',args=[order.id]))
        retrieve.assert_not_called()

    @mock.patch('stripe.checkout.Session.retrieve')
    def test_cache_read_or_lock_failure_does_not_hide_saved_order_status(self, retrieve):
        order = self.order()
        for operation in ('get', 'add'):
            with mock.patch('billing.zelda_views.cache.' + operation, side_effect=RuntimeError('Cache offline')):
                with self.assertLogs('billing.zelda_views', level='ERROR'):
                    response = self.client.get(reverse('billing:zelda_order_status',args=[order.id]))
            self.assertEqual(response.status_code,200)
            self.assertEqual(response.json()['status'],'awaiting_payment')
            self.assertTrue(response.json()['payment_check_unavailable'])
        retrieve.assert_not_called()

    @mock.patch('stripe.checkout.Session.retrieve')
    def test_cache_write_failure_does_not_hide_confirmed_payment(self, retrieve):
        order = self.order()
        retrieve.return_value = self.event(order)
        with mock.patch('billing.zelda_views.cache.set', side_effect=RuntimeError('Cache offline')):
            with self.assertLogs('billing.zelda_views', level='ERROR'), mock.patch('billing.tasks.fulfill_zelda_order.delay') as generate:
                with self.captureOnCommitCallbacks(execute=True):
                    response = self.client.get(reverse('billing:zelda_order_status',args=[order.id]))
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()['status'],'paid')
        generate.assert_called_once_with(str(order.id))

    @mock.patch('billing.fulfillment.reconcile_order', side_effect=RuntimeError('Private backend detail'))
    def test_status_failure_returns_logged_reference_without_private_details(self, reconcile):
        order = self.order(status='processing', paid_at=timezone.now())
        with self.assertLogs('billing.zelda_views', level='ERROR') as logs:
            response = self.client.get(reverse('billing:zelda_order_status',args=[order.id]))
        self.assertEqual(response.status_code,503)
        self.assertEqual(response['Retry-After'],'5')
        self.assertEqual(response['Cache-Control'],'no-store')
        reference = response.json()['reference']
        self.assertRegex(reference,r'^[a-f0-9]{12}$')
        self.assertTrue(any('reference=' + reference in line for line in logs.output))
        self.assertNotIn('Private backend detail',response.content.decode())
        self.assertNotIn('reports', response.json())

    @mock.patch('billing.zelda_views.get_object_or_404', side_effect=RuntimeError('Database unavailable'))
    def test_status_lookup_failure_also_has_an_error_reference(self, lookup):
        order = self.order()
        with self.assertLogs('billing.zelda_views', level='ERROR'):
            response = self.client.get(reverse('billing:zelda_order_status',args=[order.id]))
        self.assertEqual(response.status_code,503)
        self.assertIn('reference',response.json())

    def test_ready_report_is_saved_in_owner_library_and_not_another_library(self):
        from zelda_api.library import build_library
        doc=DocumentSource.objects.create(uploaded_by=self.user,source_entity='External Example Inc',filename='deck.txt',status='analyzed',is_external_subject=True)
        IntelligenceMemo.objects.create(document=doc,executive_summary='Grounded summary',business_model_analysis='Evidence-based model')
        order=self.order(status='processing',paid_at=timezone.now(),analysis_document=doc)
        library=build_library(self.user)
        self.assertEqual(library['purchases'][0]['status'],'Ready')
        self.assertEqual(library['purchases'][0]['reports'][0]['name'],'Zelda IC Memo')
        self.assertEqual(build_library(self.other)['purchases'],[])
        self.assertContains(self.client.get(reverse('billing:zelda_report',args=[order.id,'ic_memo'])),'Grounded summary')
        self.assertEqual(self.client.get(reverse('billing:zelda_report',args=[order.id,'truth_delta'])).status_code,404)

    def test_failed_paid_order_can_retry_without_checkout(self):
        order=self.order(status='failed',paid_at=timezone.now())
        with mock.patch('billing.tasks.fulfill_zelda_order.delay') as generate, mock.patch('stripe.checkout.Session.create') as checkout:
            with self.captureOnCommitCallbacks(execute=True):
                response=self.client.post(reverse('billing:zelda_order_retry',args=[order.id]))
        self.assertEqual(response.status_code,200);generate.assert_called_once();checkout.assert_not_called()

    def test_external_document_never_borrows_matching_uploader_profile(self):
        from matchmaking.models import Application
        from zelda_api.grounded_context import profile_for_document
        from zelda_api.entity_verification import subject_for_document
        Application.objects.create(user=self.user,company_name=self.source.source_entity)
        self.assertIsNone(profile_for_document(self.source));self.assertIsNone(subject_for_document(self.source))

    @mock.patch('zelda_api.utils.extract_text_from_file', return_value=('Readable company evidence and product information. '*30,1))
    def test_upload_stages_readable_owned_evidence_without_generation(self, extract):
        with mock.patch('billing.tasks.fulfill_zelda_order.delay') as generate:
            response=self.client.post(reverse('billing:zelda_intake'),{'company':'External Co','file':SimpleUploadedFile('deck.txt',b'readable')})
        self.assertEqual(response.status_code,201);generate.assert_not_called()
        doc=DocumentSource.objects.get(pk=response.json()['document_id'])
        self.assertEqual(doc.uploaded_by,self.user);self.assertTrue(doc.is_product_input)

    def test_unreadable_and_unsupported_uploads_are_rejected_before_purchase(self):
        with mock.patch('zelda_api.utils.extract_text_from_file',return_value=('',0)):
            response=self.client.post(reverse('billing:zelda_intake'),{'company':'External Co','file':SimpleUploadedFile('deck.txt',b'')})
        self.assertEqual(response.status_code,422)
        response=self.client.post(reverse('billing:zelda_intake'),{'company':'External Co','file':SimpleUploadedFile('script.exe',b'bad')})
        self.assertEqual(response.status_code,400);self.assertFalse(ZeldaOrder.objects.exists())

    def test_external_search_uses_shared_authority_and_never_selects_ambiguity(self):
        from zelda_api.sec_company_identity import CompanyIdentity
        with mock.patch('zelda_api.sec_company_identity.resolve_company_identity',return_value=CompanyIdentity(status='ambiguous')) as resolve:
            response=self.client.post(reverse('billing:zelda_external_search'),{'q':'Example'})
        self.assertEqual(response.json()['results'],[]);resolve.assert_called_once_with('Example')
        with mock.patch('zelda_api.sec_company_identity.resolve_company_identity',return_value=CompanyIdentity(status='found',name='Example Inc',cik='0000000001')):
            response=self.client.post(reverse('billing:zelda_external_search'),{'q':'Example Inc'})
        self.assertEqual(response.json()['results'][0]['cik'],'0000000001')

    def test_forged_company_selection_is_refused(self):
        response=self.client.post(reverse('billing:zelda_intake'),{'subject_token':'forged'})
        self.assertEqual(response.status_code,400)

    def test_public_name_choices_are_explicit_and_do_not_guess_a_registrant(self):
        choices=[{'name':'Apple Inc.','ticker':'AAPL','cik':'0000320193'}]
        with mock.patch('zelda_api.sec_company_identity.search_listed_companies',return_value=choices), mock.patch('zelda_api.sec_company_identity.resolve_company_identity') as resolve:
            response=self.client.post(reverse('billing:zelda_external_search'),{'q':'apple'})
        self.assertEqual(response.json()['results'][0]['ticker'],'AAPL')
        from django.core import signing
        selected=signing.loads(response.json()['results'][0]['token'],salt='zelda-company-selection')
        self.assertEqual(selected,{'name':'Apple Inc.','cik':'0000320193'})
        resolve.assert_not_called()

    def test_search_allows_more_than_thirty_daily_attempts_and_reports_hourly_retry(self):
        from accounts import rate_limits
        for _ in range(30):
            self.assertIsNotNone(rate_limits.reserve('zelda_company_search_user',str(self.user.pk)))
        with mock.patch('zelda_api.sec_company_identity.search_listed_companies',return_value=[{'name':'Walmart Inc.','ticker':'WMT','cik':'0000104169'}]):
            self.assertEqual(self.client.post(reverse('billing:zelda_external_search'),{'q':'Walmart'}).status_code,200)
        for _ in range(89):
            self.assertIsNotNone(rate_limits.reserve('zelda_company_search_user',str(self.user.pk)))
        response=self.client.post(reverse('billing:zelda_external_search'),{'q':'Microsoft'})
        self.assertEqual(response.status_code,429)
        self.assertIn('120 company searches per hour',response.json()['error'])
        self.assertIn('Try again in',response.json()['error'])
        self.assertGreater(int(response['Retry-After']),0)

    def test_cached_name_search_does_not_spend_the_shared_sec_budget(self):
        from accounts.models import RateLimitEvent
        cache.set('sec_listed_company_index_v1',[{'title':'Microsoft Corp','ticker':'MSFT','cik_str':789019}],3600)
        with mock.patch('zelda_api.sec_company_identity.search_listed_companies',wraps=self.real_listed_search), mock.patch('zelda_api.sec_identity._get') as get:
            response=self.client.post(reverse('billing:zelda_external_search'),{'q':'Microsoft'})
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()['results'][0]['cik'],'0000789019')
        self.assertFalse(RateLimitEvent.objects.filter(scope='identity_check_global').exists())
        get.assert_not_called()

    def test_cached_choices_work_when_shared_source_budget_is_full(self):
        from accounts import rate_limits
        for _ in range(200):
            self.assertIsNotNone(rate_limits.reserve('identity_check_global',rate_limits.GLOBAL_KEY))
        choices=[{'name':'Microsoft Corp','ticker':'MSFT','cik':'0000789019'}]
        with mock.patch('zelda_api.sec_company_identity.search_listed_companies',return_value=choices):
            response=self.client.post(reverse('billing:zelda_external_search'),{'q':'Microsoft'})
        self.assertEqual(response.status_code,200)
        response=self.client.post(reverse('billing:zelda_external_search'),{'q':'Private Co'})
        self.assertEqual(response.status_code,429)
        self.assertIn('Public-record lookups',response.json()['error'])

    def test_database_upload_failure_returns_json_without_purchase(self):
        from django.db import OperationalError
        with mock.patch('zelda_api.vector_models.DocumentSource.objects.create',side_effect=OperationalError('schema unavailable')):
            response=self.client.post(reverse('billing:zelda_intake'),{'company':'ManyChat','file':SimpleUploadedFile('deck.txt',b'Private company evidence.')})
        self.assertEqual(response.status_code,503)
        self.assertIn('could not save your document',response.json()['error'])
        self.assertFalse(ZeldaOrder.objects.exists())

    def test_library_preserves_owned_documents_when_product_storage_is_unavailable(self):
        from django.db import OperationalError
        from zelda_api.library import build_library
        with mock.patch('billing.models.ZeldaOrder.objects.filter') as orders:
            orders.return_value.only.side_effect=OperationalError('table not migrated')
            library=build_library(self.user)
        self.assertEqual(library['purchases'],[])
        self.assertIn('temporarily unavailable',library['warnings'][0])
        documents=library['sections'][0]['documents']
        self.assertEqual(documents[0]['document_id'],self.source.id)
        self.assertFalse(documents[0]['can_open_brief'])

    def test_real_pptx_upload_is_read_and_staged_without_checkout(self):
        import io
        from pptx import Presentation
        deck=Presentation()
        slide=deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text='ManyChat'
        slide.placeholders[1].text='Private company pitch deck: customer messaging platform and product overview.'
        stream=io.BytesIO();deck.save(stream)
        response=self.client.post(reverse('billing:zelda_intake'),{'company':'ManyChat','file':SimpleUploadedFile('manychat.pptx',stream.getvalue())})
        self.assertEqual(response.status_code,201)
        source=DocumentSource.objects.get(pk=response.json()['document_id'])
        self.assertEqual(source.source_entity,'ManyChat')
        self.assertEqual(source.total_pages,1)
        self.assertIn('customer messaging',source.raw_text_full)
        self.assertFalse(ZeldaOrder.objects.exists())

    def test_intelligence_memo_uses_stored_orientation_analysis(self):
        from billing.fulfillment import report_sections
        from zelda_api.principal import Principal, ORIGIN_TASK

        analysis = DocumentSource.objects.create(
            uploaded_by=self.user,
            source_entity='External Example Inc',
            filename='analysis.txt',
            document_type='pitch_deck',
            raw_text_full='Evidence',
            status='analyzed',
        )
        memo = IntelligenceMemo.objects.create(
            document=analysis,
            executive_summary='Concise company orientation.',
            business_model_analysis='Business model.',
            information_readiness='72/100 — useful evidence with gaps.',
            retrieved_context='context',
        )
        order = self.order(
            product='intelligence_memo',
            reports=['intelligence_memo'],
            amount=199,
            analysis_document=analysis,
        )
        with mock.patch('zelda_api.ic_memo.zelda_report_observations', return_value={
            'noticed': ['Revenue evidence is specific.'],
            'worth_investigating': [{'topic': 'Customer concentration', 'target': 'ic_memo'}],
        }):
            sections = report_sections(
                order,
                'intelligence_memo',
                Principal.for_user(self.user, ORIGIN_TASK, 'test'),
            )

        self.assertEqual([section['title'] for section in sections], [
            'Executive Summary', 'What Zelda Noticed', 'Worth Investigating', 'Information Readiness',
        ])
        self.assertEqual(sections[0]['text'], memo.executive_summary)
        self.assertIn('Revenue evidence is specific.', sections[1]['text'])
        self.assertIn('Customer concentration', sections[2]['text'])

    def test_existing_active_stripe_price_must_match_amount_and_currency(self):
        from .zelda_views import stripe_price
        price={'id':'price_real','unit_amount':499,'currency':'usd','type':'one_time','active':True}
        with mock.patch('stripe.Product.list') as products, mock.patch('stripe.Price.list',return_value={'data':[price]}) as prices, mock.patch('stripe.Price.retrieve',return_value=price) as retrieve:
            products.return_value.auto_paging_iter.return_value=iter([{'id':'prod_real','name':'Zelda IC Memo'}])
            self.assertEqual(stripe_price('ic_memo'),'price_real')
            self.assertEqual(products.call_args.kwargs['api_key'],'sk_test_zelda_ci_only')
            self.assertEqual(prices.call_args.kwargs['api_key'],'sk_test_zelda_ci_only')
            self.assertEqual(retrieve.call_args.kwargs['api_key'],'sk_test_zelda_ci_only')
        with mock.patch('stripe.Price.retrieve',return_value={**price,'unit_amount':500}):
            with self.assertRaises(ValueError): stripe_price('ic_memo')

    @override_settings(STRIPE_SECRET_KEY='')
    def test_missing_stripe_key_is_explained_before_any_stripe_request(self):
        with mock.patch('stripe.Product.list') as products, mock.patch('stripe.checkout.Session.create') as checkout:
            response=self.checkout()
        self.assertEqual(response.status_code,503)
        self.assertIn('Payments are not configured',response.json()['error'])
        products.assert_not_called();checkout.assert_not_called()


class ProductSearchDiscoveryTests(SimpleTestCase):
    def test_index_fetch_reserves_once_and_cached_queries_reserve_nothing(self):
        from zelda_api.sec_company_identity import search_listed_companies
        cache.clear()
        try:
            response=mock.Mock()
            response.json.return_value={'0':{'title':'Microsoft Corp','ticker':'MSFT','cik_str':789019}}
            reserve=mock.Mock()
            with mock.patch('zelda_api.sec_identity._get',return_value=response) as get:
                self.assertEqual(search_listed_companies('Microsoft',before_fetch=reserve)[0]['ticker'],'MSFT')
                self.assertEqual(search_listed_companies('MSFT',before_fetch=reserve)[0]['ticker'],'MSFT')
            reserve.assert_called_once();get.assert_called_once()
        finally: cache.clear()

    def test_company_names_legal_suffixes_and_tickers_find_the_same_choices(self):
        from zelda_api.sec_company_identity import listed_company_matches
        rows=[{'title':'Apple Inc.','ticker':'AAPL','cik_str':320193},
              {'title':'Walmart Inc.','ticker':'WMT','cik_str':104169},
              {'title':'NIKE, Inc.','ticker':'NKE','cik_str':320187},
              {'title':'Crocs, Inc.','ticker':'CROX','cik_str':1334036},
              {'title':'Apple Hospitality REIT, Inc.','ticker':'APLE','cik_str':1418121}]
        for query,cik in [('apple','0000320193'),('Apple Inc','0000320193'),('aapl','0000320193'),
                          ('walmart','0000104169'),('WMT','0000104169'),('nike','0000320187'),('crocs','0001334036')]:
            with self.subTest(query=query):
                self.assertEqual([row['cik'] for row in listed_company_matches(query,rows)],[cik])
        self.assertEqual(listed_company_matches('Ben and Jerry',rows),[])
        self.assertEqual(listed_company_matches('appleton',rows),[])

    def test_oversized_product_upload_is_json_and_refused_before_body_processing(self):
        from shared_utils.upload_limits import UploadSizeLimitMiddleware, MB
        downstream=mock.Mock()
        request=RequestFactory().post(reverse('billing:zelda_intake'))
        request.META['CONTENT_LENGTH']=str(26*MB)
        response=UploadSizeLimitMiddleware(downstream)(request)
        self.assertEqual(response.status_code,413)
        self.assertIn('25 MB',json.loads(response.content)['error'])
        downstream.assert_not_called()

    def test_source_failure_is_not_cached_as_a_missing_company(self):
        from zelda_api.sec_company_identity import search_listed_companies
        from zelda_api.sec_identity import SecUnavailable
        cache.clear()
        try:
            with mock.patch('zelda_api.sec_identity._get',side_effect=SecUnavailable('unavailable')) as get:
                for _ in range(2):
                    with self.assertRaises(SecUnavailable): search_listed_companies('Apple')
                self.assertEqual(get.call_count,2)
        finally: cache.clear()
